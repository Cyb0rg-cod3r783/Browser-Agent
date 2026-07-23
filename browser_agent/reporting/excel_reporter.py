"""
reporting/excel_reporter.py — Generates a styled Excel test report.
"""
import os
import json
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from schema import TestResult

def format_steps(test_case_steps_json, elements_map):
    steps = json.loads(test_case_steps_json) if isinstance(test_case_steps_json, str) else test_case_steps_json
    if not steps:
        return "1. Go to login page\n2. Perform happy path steps"
        
    lines = ["1. Go to login page: https://test5.squad1.tech/login?RetUrl=/settings/usermanagement/employee"]
    for i, s in enumerate(steps, 2):
        action = s.get('action', '').capitalize()
        el_id = s.get('element_id')
        label = elements_map.get(el_id, "element")
        val = s.get('value')
        
        if val == "":
            lines.append(f"{i}. {action} '{label}' with empty value")
        elif val is not None and val != "None":
            lines.append(f"{i}. {action} '{label}' with value '{val}'")
        else:
            lines.append(f"{i}. {action} '{label}'")
    return "\n".join(lines)

def get_test_data_smart(test_name, test_case_steps_json, happy_path_steps, elements_map):
    if "Happy Path" in test_name:
        return "NA"
        
    steps = json.loads(test_case_steps_json) if isinstance(test_case_steps_json, str) else test_case_steps_json
    if not steps:
        return "NA"
        
    for s, h_s in zip(steps, happy_path_steps):
        s_val = s.get('value')
        h_val = h_s.get('value')
        s_el = s.get('element_id')
        h_el = h_s.get('element_id')
        
        if s_val != h_val:
            label = elements_map.get(s_el, "Field")
            if s_val == "":
                return f"{label}: [Empty]"
            return f"{label}: \"{s_val}\""
        elif s_el != h_el:
            label = elements_map.get(s_el, "Field")
            h_label = elements_map.get(h_el, "Field")
            return f"{label}: Selected \"{label}\" (was \"{h_label}\")"
            
    field_name_part = test_name.split("-")[0].strip().lower()
    for s in steps:
        s_el = s.get('element_id')
        label = elements_map.get(s_el, "").lower()
        if field_name_part in label or label in field_name_part:
            val = s.get('value')
            if val == "":
                return f"{elements_map.get(s_el)}: [Empty]"
            return f"{elements_map.get(s_el)}: \"{val}\""
            
    return "NA"

def format_expected(assertion_results, test_name):
    assertions = json.loads(assertion_results) if isinstance(assertion_results, str) else assertion_results
    if not assertions:
        return "1. Successful form submission"
        
    lines = []
    for i, a in enumerate(assertions, 1):
        a_type = a.get('type')
        expected = a.get('expected')
        if a_type == "url_contains":
            lines.append(f"{i}. The URL should contain '{expected}'")
        elif a_type == "element_visible":
            lines.append(f"{i}. The element/message '{expected}' should be visible")
        elif a_type == "text_equals":
            lines.append(f"{i}. The validation error message '{expected}' should be displayed")
        else:
            lines.append(f"{i}. Expect {a_type}: '{expected}'")
    return "\n".join(lines)

def format_actual(assertion_results, status, error_detail):
    if status != "passed":
        return f"Failed. {error_detail or 'Mismatch'}"
    assertions = json.loads(assertion_results) if isinstance(assertion_results, str) else assertion_results
    if not assertions:
        return "As Expected"
    lines = []
    for a in assertions:
        actual = a.get('actual')
        if actual is None:
            actual = "Passed"
        lines.append(str(actual))
    if len(lines) == 1:
        return lines[0]
    return "\n".join(f"{i}. {line}" for i, line in enumerate(lines, 1))

async def generate_excel_report(
    run_id: str,
    results: list[TestResult],
    db,
    output_path: str,
    app_name: str
) -> None:
    """
    Generate a styled Excel test report.
    """
    # 1. Fetch element mapping from DB
    conn = sqlite3.connect(db.db_path)
    cur = conn.cursor()
    
    cur.execute("SELECT id FROM applications WHERE name=?", (app_name,))
    app_row = cur.fetchone()
    if not app_row:
        conn.close()
        return
    app_id = app_row[0]
    
    cur.execute("SELECT id FROM app_versions WHERE app_id=? AND is_active=1", (app_id,))
    version_row = cur.fetchone()
    if not version_row:
        conn.close()
        return
    version_id = version_row[0]
    
    cur.execute("SELECT id, semantic_label FROM elements WHERE page_id IN (SELECT id FROM pages WHERE version_id=?)", (version_id,))
    elements_map = {r[0]: r[1] for r in cur.fetchall()}
    
    # Fetch test case steps from DB for step formatting
    cur.execute("SELECT id, steps FROM test_cases WHERE flow_id IN (SELECT id FROM user_flows WHERE version_id=?)", (version_id,))
    tc_steps_map = {r[0]: r[1] for r in cur.fetchall()}
    conn.close()

    # 2. Sort results: happy path first, then alphabetically
    sorted_results = []
    happy_paths = [r for r in results if "Happy Path" in r.test_name]
    others = [r for r in results if "Happy Path" not in r.test_name]
    others.sort(key=lambda x: x.test_name)
    
    sorted_results.extend(happy_paths)
    sorted_results.extend(others)

    # Load Happy Path steps
    happy_path_steps = []
    if happy_paths:
        tc_steps_json = tc_steps_map.get(happy_paths[0].test_case_id)
        if tc_steps_json:
            happy_path_steps = json.loads(tc_steps_json)

    # 3. Create openpyxl workbook
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"
    sheet.views.sheetView[0].showGridLines = True

    # Column widths
    col_widths = {
        'A': 13.0, 'B': 13.25, 'C': 12.625, 'D': 11.875, 'E': 18.625,
        'F': 40.375, 'G': 13.125, 'H': 32.625, 'I': 13.125, 'J': 17.25, 'K': 13.0
    }
    for col, w in col_widths.items():
        sheet.column_dimensions[col].width = w

    sheet['F2'] = 'in'

    # Headers
    headers = [
        'Module', 'Sub-Module', 'Test Case Id', 'Test Case Scenario',
        'Test Case Steps', 'Test Data', 'Expected Result', 'Actual Result',
        'Status (Pass/Fail)'
    ]
    
    yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    thin_border_side = Side(border_style="thin", color="A0A0A0")
    thin_border = Border(
        top=thin_border_side, bottom=thin_border_side,
        left=thin_border_side, right=thin_border_side
    )

    sheet.row_dimensions[3].height = 25.0
    for i, h in enumerate(headers, 2):
        cell = sheet.cell(row=3, column=i)
        cell.value = h
        cell.fill = yellow_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # Data rows
    data_font = Font(name="Calibri", size=11, bold=False, color="000000")
    data_align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    data_align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    start_row = 4
    for idx, r in enumerate(sorted_results):
        row_num = start_row + idx
        sheet.row_dimensions[row_num].height = 90.0
        
        tc_steps_json = tc_steps_map.get(r.test_case_id, "[]")
        tc_id = f"TC_EMP_{idx+1:02d}"
        scenario = f"To verify {r.test_name}"
        steps = format_steps(tc_steps_json, elements_map)
        test_data = get_test_data_smart(r.test_name, tc_steps_json, happy_path_steps, elements_map)
        
        expected = format_expected(r.assertion_results, r.test_name)
        actual = format_actual(r.assertion_results, r.status, r.error_detail)
        status_text = "PASS" if r.status == "passed" else "FAIL"
        
        if row_num == start_row:
            sheet.cell(row=row_num, column=2, value="Employee Management")
            
        sheet.cell(row=row_num, column=3, value=None)
        sheet.cell(row=row_num, column=4, value=tc_id)
        sheet.cell(row=row_num, column=5, value=scenario)
        sheet.cell(row=row_num, column=6, value=steps)
        sheet.cell(row=row_num, column=7, value=test_data)
        sheet.cell(row=row_num, column=8, value=expected)
        sheet.cell(row=row_num, column=9, value=actual)
        sheet.cell(row=row_num, column=10, value=status_text)
        
        for col_idx in range(2, 11):
            cell = sheet.cell(row=row_num, column=col_idx)
            cell.border = thin_border
            if col_idx in (2, 4, 10):
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left
                
            if col_idx == 10:
                if status_text == "PASS":
                    cell.font = Font(name="Calibri", size=11, bold=True, color="008000")
                else:
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FF0000")
            else:
                cell.font = data_font

    # Merge Module B
    end_row = start_row + len(sorted_results) - 1
    sheet.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
    
    for r in range(start_row, end_row + 1):
        cell = sheet.cell(row=r, column=2)
        cell.border = thin_border
        cell.alignment = data_align_center
        cell.font = data_font

    # Save to output path
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
